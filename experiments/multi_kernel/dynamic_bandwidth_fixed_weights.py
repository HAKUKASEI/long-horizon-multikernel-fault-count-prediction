import os
import glob
import gc
import tempfile
import traceback
import faulthandler
import multiprocessing as mp
from concurrent.futures import ProcessPoolExecutor, as_completed
from multiprocessing import freeze_support

# Better defaults for server-side multi-process execution:
# avoid nested oversubscription from BLAS/OpenMP libraries.
os.environ.setdefault("OMP_NUM_THREADS", "1")
os.environ.setdefault("MKL_NUM_THREADS", "1")
os.environ.setdefault("OPENBLAS_NUM_THREADS", "1")
os.environ.setdefault("NUMEXPR_NUM_THREADS", "1")

# Make matplotlib / numba cache writable on servers.
os.environ.setdefault("MPLCONFIGDIR", os.path.join(tempfile.gettempdir(), "mplconfig"))
os.environ.setdefault("NUMBA_CACHE_DIR", os.path.join(tempfile.gettempdir(), "numba_cache"))
os.makedirs(os.environ["MPLCONFIGDIR"], exist_ok=True)
os.makedirs(os.environ["NUMBA_CACHE_DIR"], exist_ok=True)

try:
    faulthandler.enable()
except Exception:
    pass

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy.optimize import minimize_scalar
from numba import njit

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

try:
    from classicUtils import SRMtool
except Exception:
    SRMtool = None


# ============================================================
# 1) Environment helpers
# ============================================================
def setup_matplotlib_fonts():
    plt.rcParams["axes.unicode_minus"] = False


def safe_send_notice_mail(subject, body, to_email):
    if SRMtool is None:
        print(f"[Mail skipped] {subject} -> classicUtils.SRMtool not available")
        return
    try:
        SRMtool.sendNoticeMail(subject, body, to_email)
    except Exception as e:
        print(f"[Mail warning] {e}")


def recommend_max_workers(requested=None):
    """
    Safer default for servers:
    use at most half CPU cores, capped at 8 workers.
    """
    cpu = os.cpu_count() or 2
    safe_default = max(1, min(cpu // 2, 8))

    if requested is None:
        return safe_default

    requested = int(requested)
    return max(1, min(requested, cpu))


# ============================================================
# 2) Kernel library (public Python wrappers)
# ============================================================
class Kernels:
    @staticmethod
    def Gaussian(t, h):
        x = t / h
        return (1.0 / h) * (1.0 / np.sqrt(2.0 * np.pi)) * np.exp(-0.5 * x**2)

    @staticmethod
    def Triangular(t, h):
        x = t / h
        return np.where(np.abs(x) <= 1.0, (1.0 - np.abs(x)) / h, 0.0)

    @staticmethod
    def Biweight(t, h):
        x = t / h
        return np.where(np.abs(x) <= 1.0, (15.0 / 16.0) * ((1.0 - x**2) ** 2) / h, 0.0)

    @staticmethod
    def Epanechnikov(t, h):
        x = t / h
        return np.where(np.abs(x) <= 1.0, 0.75 * (1.0 - x**2) / h, 0.0)

    @staticmethod
    def Uniform(t, h):
        x = t / h
        return np.where(np.abs(x) <= 1.0, 0.5 / h, 0.0)


KERNEL_IDS = {
    "Gaussian": 0,
    "Epanechnikov": 1,
    "Triangular": 2,
    "Biweight": 3,
    "Uniform": 4,
}


def get_kernel_map():
    return {
        "Gaussian": Kernels.Gaussian,
        "Epanechnikov": Kernels.Epanechnikov,
        "Triangular": Kernels.Triangular,
        "Biweight": Kernels.Biweight,
        "Uniform": Kernels.Uniform,
    }


def _kernel_id_from_func(kernel_func):
    name = getattr(kernel_func, "__name__", None)
    if name not in KERNEL_IDS:
        raise ValueError(f"Unsupported kernel function: {name}")
    return KERNEL_IDS[name]


# ============================================================
# 3) NUMBA cores
# ============================================================
@njit(cache=True)
def _kernel_scalar_numba(t, h, kernel_id):
    x = t / h
    ax = abs(x)

    if kernel_id == 0:  # Gaussian
        return (1.0 / h) * (1.0 / np.sqrt(2.0 * np.pi)) * np.exp(-0.5 * x * x)
    elif kernel_id == 1:  # Epanechnikov
        if ax <= 1.0:
            return 0.75 * (1.0 - x * x) / h
        return 0.0
    elif kernel_id == 2:  # Triangular
        if ax <= 1.0:
            return (1.0 - ax) / h
        return 0.0
    elif kernel_id == 3:  # Biweight
        if ax <= 1.0:
            return (15.0 / 16.0) * ((1.0 - x * x) ** 2) / h
        return 0.0
    else:  # Uniform
        if ax <= 1.0:
            return 0.5 / h
        return 0.0


@njit(cache=True)
def _trapz_numba(y, x):
    n = len(x)
    if n <= 1:
        return 0.0

    s = 0.0
    for i in range(n - 1):
        s += 0.5 * (y[i] + y[i + 1]) * (x[i + 1] - x[i])
    return s


@njit(cache=True)
def _cumulative_trapz_numba(y, x):
    n = len(x)
    out = np.zeros(n, dtype=np.float64)
    if n <= 1:
        return out

    s = 0.0
    for i in range(n - 1):
        s += 0.5 * (y[i] + y[i + 1]) * (x[i + 1] - x[i])
        out[i + 1] = s
    return out


@njit(cache=True)
def _get_intensity_scalar_numba(u_grid, U_events, h, kernel_id):
    n_u = len(u_grid)
    n_evt = len(U_events)
    out = np.zeros(n_u, dtype=np.float64)

    for i in range(n_u):
        ui = u_grid[i]
        s = 0.0
        for j in range(n_evt):
            s += _kernel_scalar_numba(ui - U_events[j], h, kernel_id)
        out[i] = s
    return out


@njit(cache=True)
def _boundary_corr_scalar_numba(u_grid, h, kernel_id, window_left, window_right, v_grid_size):
    n_u = len(u_grid)
    out = np.ones(n_u, dtype=np.float64)

    if v_grid_size <= 1 or window_right <= window_left:
        return out

    v_grid = np.linspace(window_left, window_right, v_grid_size)

    for i in range(n_u):
        ui = u_grid[i]
        vals = np.empty(v_grid_size, dtype=np.float64)
        for k in range(v_grid_size):
            vals[k] = _kernel_scalar_numba(ui - v_grid[k], h, kernel_id)
        p = _trapz_numba(vals, v_grid)
        if p < 1e-12:
            p = 1e-12
        out[i] = p

    return out


@njit(cache=True)
def _get_intensity_corrected_scalar_numba(u_grid, U_events, h, kernel_id, window_left, window_right, v_grid_size):
    n_u = len(u_grid)
    n_evt = len(U_events)
    out = np.zeros(n_u, dtype=np.float64)

    if n_evt == 0:
        return out

    p_u = _boundary_corr_scalar_numba(
        u_grid, h, kernel_id, window_left, window_right, v_grid_size
    )

    for i in range(n_u):
        ui = u_grid[i]
        s = 0.0
        for j in range(n_evt):
            s += _kernel_scalar_numba(ui - U_events[j], h, kernel_id)
        out[i] = s / p_u[i]

    return out


@njit(cache=True)
def _loo_row_sums_scalar_numba(U_events, h, kernel_id):
    n = len(U_events)
    out = np.zeros(n, dtype=np.float64)

    for i in range(n):
        ui = U_events[i]
        s = 0.0
        for j in range(n):
            if i != j:
                s += _kernel_scalar_numba(ui - U_events[j], h, kernel_id)
        out[i] = s

    return out


@njit(cache=True)
def _weighted_pred_intensity_from_khat_scalar_numba(x_grid, X_events, h, k_hat_at_events, kernel_id):
    n_grid = len(x_grid)
    n_evt = len(X_events)
    out = np.zeros(n_grid, dtype=np.float64)

    if n_evt == 0:
        return out

    for i in range(n_grid):
        xi = x_grid[i]

        # Gaussian: numerically stable normalized local weights.
        # The common Gaussian factor cancels in the ratio, so this
        # preserves the original formula while avoiding far-tail underflow.
        if kernel_id == 0:
            max_logw = -1.0e308

            for j in range(n_evt):
                z = (xi - X_events[j]) / h
                logw = -0.5 * z * z
                if logw > max_logw:
                    max_logw = logw

            numer = 0.0
            denom = 0.0
            for j in range(n_evt):
                z = (xi - X_events[j]) / h
                logw = -0.5 * z * z
                w = np.exp(logw - max_logw)
                denom += w
                numer += w * k_hat_at_events[j]

            if denom <= 0.0:
                out[i] = 0.0
            else:
                val = numer / denom
                out[i] = val if val > 0.0 else 0.0

        # Compact-support kernels: a zero denominator outside support is genuine.
        else:
            numer = 0.0
            denom = 0.0
            for j in range(n_evt):
                kval = _kernel_scalar_numba(xi - X_events[j], h, kernel_id)
                denom += kval
                numer += kval * k_hat_at_events[j]

            if denom < 1e-12:
                out[i] = 0.0
            else:
                val = numer / denom
                out[i] = val if val > 0.0 else 0.0

    return out


# ============================================================
# 4) Base intensity estimators on normalized scale
# ============================================================
def get_intensity(u_grid, U_events, h, kernel_func):
    u_grid = np.asarray(u_grid, dtype=np.float64)
    U_events = np.asarray(U_events, dtype=np.float64)
    h = float(h)
    kernel_id = _kernel_id_from_func(kernel_func)
    return _get_intensity_scalar_numba(u_grid, U_events, h, kernel_id)


def correction_factor(u_grid, h, kernel_func, x_cur, v_grid_size=1000):
    u_grid = np.asarray(u_grid, dtype=np.float64)
    h = float(h)
    x_cur = float(x_cur)
    kernel_id = _kernel_id_from_func(kernel_func)
    return _boundary_corr_scalar_numba(
        u_grid,
        h,
        kernel_id,
        0.0,
        x_cur,
        int(v_grid_size),
    )


def get_intensity_corrected(u_grid, U_events, h, kernel_func, x_cur, v_grid_size=1000):
    u_grid = np.asarray(u_grid, dtype=np.float64)
    U_events = np.asarray(U_events, dtype=np.float64)
    h = float(h)
    x_cur = float(x_cur)
    kernel_id = _kernel_id_from_func(kernel_func)
    return _get_intensity_corrected_scalar_numba(
        u_grid,
        U_events,
        h,
        kernel_id,
        0.0,
        x_cur,
        int(v_grid_size),
    )


def get_intensity_by_type(
    x_grid,
    X_events,
    h,
    kernel_func,
    estimator_type="ordinary",
    x_cur=None,
    v_grid_size=1000,
):
    if estimator_type == "ordinary":
        return get_intensity(x_grid, X_events, h, kernel_func)
    elif estimator_type == "corrected":
        if x_cur is None:
            x_cur = float(np.max(X_events)) if len(X_events) else 1.0
        return get_intensity_corrected(
            x_grid,
            X_events,
            h,
            kernel_func,
            x_cur=x_cur,
            v_grid_size=v_grid_size,
        )
    else:
        raise ValueError(f"Unknown estimator_type: {estimator_type}")


def evaluate_intensity_at_events(
    U_events_eval,
    U_events_train,
    h,
    kernel_func,
    corrected=False,
    x_cur=None,
    v_grid_size=1000,
):
    U_events_eval = np.asarray(U_events_eval, dtype=np.float64)
    U_events_train = np.asarray(U_events_train, dtype=np.float64)
    h = float(h)

    if not corrected:
        lam = get_intensity(U_events_eval, U_events_train, h, kernel_func)
        return np.maximum(lam, 1e-12)

    if x_cur is None:
        x_cur = float(np.max(U_events_train)) if len(U_events_train) else 1.0

    lam = get_intensity_corrected(
        U_events_eval,
        U_events_train,
        h,
        kernel_func,
        x_cur=float(x_cur),
        v_grid_size=v_grid_size,
    )
    return np.maximum(lam, 1e-12)


def evaluate_intensity_loo_at_events_ordinary(U_events, h, kernel_func):
    U_events = np.asarray(U_events, dtype=np.float64)
    h = float(h)
    kernel_id = _kernel_id_from_func(kernel_func)
    lam_loo = _loo_row_sums_scalar_numba(U_events, h, kernel_id)
    return np.maximum(lam_loo, 1e-12)


def evaluate_intensity_loo_at_events_corrected(
    U_events,
    h,
    kernel_func,
    x_cur,
    v_grid_size=1000,
):
    U_events = np.asarray(U_events, dtype=np.float64)
    h = float(h)
    x_cur = float(x_cur)
    kernel_id = _kernel_id_from_func(kernel_func)

    row_sums = _loo_row_sums_scalar_numba(U_events, h, kernel_id)
    p_u = _boundary_corr_scalar_numba(U_events, h, kernel_id, 0.0, x_cur, int(v_grid_size))
    lam_loo = row_sums / np.maximum(p_u, 1e-12)
    return np.maximum(lam_loo, 1e-12)


# ============================================================
# 5) Campbell score + bandwidth scores
# ============================================================
def compute_campbell_score(
    U_events,
    h,
    kernel_func,
    corrected=False,
    x_cur=None,
    v_grid_size=1000,
):
    U_events = np.asarray(U_events, dtype=np.float64)

    if x_cur is None:
        x_cur = float(np.max(U_events)) if len(U_events) else 1.0
    x_cur = float(x_cur)

    lam_at_events = evaluate_intensity_at_events(
        U_events_eval=U_events,
        U_events_train=U_events,
        h=h,
        kernel_func=kernel_func,
        corrected=corrected,
        x_cur=x_cur,
        v_grid_size=v_grid_size,
    )

    t_kappa = np.sum(1.0 / np.maximum(lam_at_events, 1e-12))
    score = (t_kappa - x_cur) ** 2
    return float(score)


def evaluate_score_values(
    X_obs,
    h,
    x_grid_obs,
    kernel_func,
    estimator_type="ordinary",
    v_grid_size=1000,
):
    X_obs = np.asarray(X_obs, dtype=np.float64)
    x_grid_obs = np.asarray(x_grid_obs, dtype=np.float64)
    h = float(h)

    x_cur = float(np.max(X_obs)) if len(X_obs) else 1.0
    kernel_id = _kernel_id_from_func(kernel_func)

    if estimator_type == "ordinary":
        lam = _get_intensity_scalar_numba(x_grid_obs, X_obs, h, kernel_id)
        lam_at_events = _get_intensity_scalar_numba(X_obs, X_obs, h, kernel_id)
        lam_loo = _loo_row_sums_scalar_numba(X_obs, h, kernel_id)
        corrected = False
    elif estimator_type == "corrected":
        lam = _get_intensity_corrected_scalar_numba(
            x_grid_obs,
            X_obs,
            h,
            kernel_id,
            0.0,
            x_cur,
            int(v_grid_size),
        )
        lam_at_events = _get_intensity_corrected_scalar_numba(
            X_obs,
            X_obs,
            h,
            kernel_id,
            0.0,
            x_cur,
            int(v_grid_size),
        )
        row_sums = _loo_row_sums_scalar_numba(X_obs, h, kernel_id)
        p_u = _boundary_corr_scalar_numba(X_obs, h, kernel_id, 0.0, x_cur, int(v_grid_size))
        lam_loo = row_sums / np.maximum(p_u, 1e-12)
        corrected = True
    else:
        raise ValueError(f"Unknown estimator_type: {estimator_type}")

    integral = _trapz_numba(np.asarray(lam, dtype=np.float64), x_grid_obs)

    s_ml = np.sum(np.log(np.maximum(lam_at_events, 1e-12))) - integral
    s_llcv = np.sum(np.log(np.maximum(lam_loo, 1e-12))) - integral

    term1 = _trapz_numba(np.asarray(lam * lam, dtype=np.float64), x_grid_obs)
    term2 = 2.0 * np.sum(lam_loo)
    s_lscv = term1 - term2

    s_campbell = compute_campbell_score(
        U_events=X_obs,
        h=h,
        kernel_func=kernel_func,
        corrected=corrected,
        x_cur=x_cur,
        v_grid_size=v_grid_size,
    )
    return float(s_ml), float(s_llcv), float(s_lscv), float(s_campbell)


def compute_scores(
    X_obs,
    h_cands,
    x_grid_obs,
    kernel_func,
    estimator_type="ordinary",
    v_grid_size=1000,
):
    h_cands = np.asarray(h_cands, dtype=np.float64)

    s_ml, s_llcv, s_lscv, s_campbell = [], [], [], []

    for h in h_cands:
        ml, llcv, lscv, campbell = evaluate_score_values(
            X_obs,
            float(h),
            x_grid_obs,
            kernel_func,
            estimator_type=estimator_type,
            v_grid_size=v_grid_size,
        )
        s_ml.append(ml)
        s_llcv.append(llcv)
        s_lscv.append(lscv)
        s_campbell.append(campbell)

    return (
        np.asarray(s_ml, dtype=np.float64),
        np.asarray(s_llcv, dtype=np.float64),
        np.asarray(s_lscv, dtype=np.float64),
        np.asarray(s_campbell, dtype=np.float64),
    )


# ============================================================
# 6) Maximum likelihood weights (Log-Sum-Exp softmax)
# ============================================================
def max_likelihood_weights(log_likelihoods):
    ll = np.asarray(log_likelihoods, dtype=np.float64)
    ll_max = np.max(ll)
    exp_terms = np.exp(ll - ll_max)
    return exp_terms / np.sum(exp_terms)


def loglik_mle_on_window(
    X_events,
    h,
    kernel_func,
    x_grid_base,
    x_cur,
    estimator_type="ordinary",
    v_grid_size=1000,
):
    X_events = np.asarray(X_events, dtype=np.float64)
    x_cur = float(x_cur)
    h = float(h)

    if x_cur <= 0.0:
        return 0.0

    X_use = X_events[X_events <= x_cur + 1e-12]
    if len(X_use) == 0:
        return 0.0

    base_n = len(x_grid_base) if x_grid_base is not None else 2000
    n_grid = max(200, int(np.ceil(base_n * x_cur)))
    x_grid = np.linspace(0.0, x_cur, n_grid, dtype=np.float64)

    lam = get_intensity_by_type(
        x_grid,
        X_use,
        h,
        kernel_func,
        estimator_type=estimator_type,
        x_cur=x_cur,
        v_grid_size=v_grid_size,
    )

    lam_at_events = evaluate_intensity_at_events(
        U_events_eval=X_use,
        U_events_train=X_use,
        h=h,
        kernel_func=kernel_func,
        corrected=(estimator_type == "corrected"),
        x_cur=x_cur,
        v_grid_size=v_grid_size,
    )

    integral = _trapz_numba(np.asarray(lam, dtype=np.float64), x_grid)
    return float(np.sum(np.log(np.maximum(lam_at_events, 1e-12))) - integral)


# ============================================================
# 7) Prediction intensity
# ============================================================
def weighted_pred_intensity(
    x_grid,
    X_events,
    h,
    kernel_func,
    estimator_type="ordinary",
    v_grid_size=1000,
):
    X_events = np.asarray(X_events, dtype=np.float64)
    x_grid = np.asarray(x_grid, dtype=np.float64)
    h = float(h)

    if len(X_events) == 0:
        return np.zeros_like(x_grid, dtype=np.float64)

    x_cur = float(np.max(X_events))

    if estimator_type == "ordinary":
        k_hat_at_events = evaluate_intensity_at_events(
            U_events_eval=X_events,
            U_events_train=X_events,
            h=h,
            kernel_func=kernel_func,
            corrected=False,
            x_cur=x_cur,
            v_grid_size=v_grid_size,
        )
    elif estimator_type == "corrected":
        k_hat_at_events = evaluate_intensity_at_events(
            U_events_eval=X_events,
            U_events_train=X_events,
            h=h,
            kernel_func=kernel_func,
            corrected=True,
            x_cur=x_cur,
            v_grid_size=v_grid_size,
        )
    else:
        raise ValueError(f"Unknown estimator_type: {estimator_type}")

    kernel_id = _kernel_id_from_func(kernel_func)
    return _weighted_pred_intensity_from_khat_scalar_numba(
        x_grid,
        X_events,
        h,
        np.asarray(k_hat_at_events, dtype=np.float64),
        kernel_id,
    )


def ensemble_pred_intensity(
    x_grid,
    X_events,
    kernel_configs,
    weights,
    estimator_type="ordinary",
    w_eps=0.0,
    v_grid_size=1000,
):
    x_grid = np.asarray(x_grid, dtype=np.float64)
    total = np.zeros_like(x_grid, dtype=np.float64)

    for cfg, w in zip(kernel_configs, weights):
        if w <= w_eps:
            continue
        kp = weighted_pred_intensity(
            x_grid,
            X_events,
            cfg["h"],
            cfg["func"],
            estimator_type=estimator_type,
            v_grid_size=v_grid_size,
        )
        total += float(w) * kp

    return np.maximum(total, 0.0)


# ============================================================
# 8) Recursive prediction support
# ============================================================
def cumulative_trapz(y, x):
    y = np.asarray(y, dtype=np.float64)
    x = np.asarray(x, dtype=np.float64)
    return _cumulative_trapz_numba(y, x)


def find_next_event_time_ensemble(
    x_cur,
    X_events,
    kernel_configs,
    weights,
    x_end,
    estimator_type="ordinary",
    base_ngrid=600,
    max_expand=60,
    v_grid_size=1000,
):
    x_cur = float(x_cur)
    x_end = float(x_end)

    if x_cur >= x_end:
        return x_end, False

    hs = [float(cfg["h"]) for cfg in kernel_configs]
    h_ref = float(np.median(hs)) if len(hs) else 1e-3

    initial_span = max(2.0 * h_ref, 1e-3)
    total_span = max(x_end - x_cur, 0.0)
    upper = min(x_end, x_cur + initial_span)

    # Ensure that the search can always reach x_end.
    # max_expand is only a minimum safety budget.
    if total_span > initial_span and initial_span > 0.0:
        needed_expand = int(np.ceil(np.log2(total_span / initial_span))) + 2
    else:
        needed_expand = 1
    effective_max_expand = max(int(max_expand), needed_expand)

    for _ in range(effective_max_expand):
        if upper <= x_cur + 1e-12:
            upper = min(x_end, x_cur + 1e-3)

        # Increase integration resolution as the searched interval expands.
        span_ratio = max((upper - x_cur) / initial_span, 1.0)
        n_grid = int(min(5000, max(base_ngrid, np.ceil(base_ngrid * span_ratio))))

        x_grid = np.linspace(x_cur, upper, n_grid, dtype=np.float64)
        kp_ens = ensemble_pred_intensity(
            x_grid,
            X_events,
            kernel_configs,
            weights,
            estimator_type=estimator_type,
            v_grid_size=v_grid_size,
        )
        cum = cumulative_trapz(kp_ens, x_grid)
        I = float(cum[-1])

        if I >= 1.0:
            x_next = np.interp(1.0, cum, x_grid)
            return float(x_next), True

        if np.isclose(upper, x_end) or upper >= x_end:
            return float(x_end), False

        new_upper = min(x_end, x_cur + 2.0 * (upper - x_cur))
        if new_upper <= upper + 1e-12:
            upper = x_end
        else:
            upper = new_upper

    # Final fallback: explicitly check the full remaining interval.
    x_grid = np.linspace(x_cur, x_end, 5000, dtype=np.float64)
    kp_ens = ensemble_pred_intensity(
        x_grid,
        X_events,
        kernel_configs,
        weights,
        estimator_type=estimator_type,
        v_grid_size=v_grid_size,
    )
    cum = cumulative_trapz(kp_ens, x_grid)
    I = float(cum[-1])

    if I >= 1.0:
        x_next = np.interp(1.0, cum, x_grid)
        return float(x_next), True

    return float(x_end), False


# ============================================================
# 9) PAE metrics
# ============================================================
def compute_pae_metrics_piecewise_linear(X_obs, preds, X_val, x_end):
    X_obs = np.asarray(X_obs, dtype=np.float64)
    preds = np.asarray(preds, dtype=np.float64)
    X_val = np.asarray(X_val, dtype=np.float64)
    x_end = float(x_end)

    if len(X_val) == 0:
        return np.nan, np.nan, None, None, None, None, None

    if len(preds) > 0:
        X_pred_all = np.concatenate([X_obs, preds])
    else:
        X_pred_all = X_obs.copy()
    X_pred_all = np.sort(X_pred_all)

    x_knots = np.concatenate(([0.0], X_pred_all, [x_end]))
    y_knots = np.concatenate(
        ([0.0], np.arange(1, len(X_pred_all) + 1, dtype=np.float64), [float(len(X_pred_all))])
    )

    y_true_at_val = len(X_obs) + np.arange(1, len(X_val) + 1, dtype=np.float64)
    a_at_val = np.interp(X_val, x_knots, y_knots)

    AE_curve = np.abs(y_true_at_val - a_at_val)
    PAE_MAE = float(np.mean(AE_curve))
    PAE_RMSE = float(np.sqrt(np.mean(AE_curve**2)))

    return PAE_MAE, PAE_RMSE, x_knots, y_knots, y_true_at_val, a_at_val, AE_curve


# ============================================================
# 10) Excel summary output
# ============================================================
def write_summary_excel(excel_path, rows_main, rows_kernel, rows_kernel_step):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Summary"
    if len(rows_main) == 0:
        ws1.append(["No results"])
    else:
        headers = list(rows_main[0].keys())
        ws1.append(headers)
        for r in rows_main:
            ws1.append([r.get(h) for h in headers])

        for j, h in enumerate(headers, start=1):
            max_len = max(len(str(h)), max(len(str(r.get(h))) for r in rows_main))
            ws1.column_dimensions[get_column_letter(j)].width = min(45, max(10, max_len + 2))
        ws1.freeze_panes = "A2"

    ws2 = wb.create_sheet("KernelWeights")
    if len(rows_kernel) == 0:
        ws2.append(["No kernel-weight results"])
    else:
        headers2 = list(rows_kernel[0].keys())
        ws2.append(headers2)
        for r in rows_kernel:
            ws2.append([r.get(h) for h in headers2])

        for j, h in enumerate(headers2, start=1):
            max_len = max(len(str(h)), max(len(str(r.get(h))) for r in rows_kernel))
            ws2.column_dimensions[get_column_letter(j)].width = min(45, max(10, max_len + 2))
        ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("KernelWeightsByStep")
    if len(rows_kernel_step) == 0:
        ws3.append(["No per-step kernel-weight results"])
    else:
        headers3 = list(rows_kernel_step[0].keys())
        ws3.append(headers3)
        for r in rows_kernel_step:
            ws3.append([r.get(h) for h in headers3])

        for j, h in enumerate(headers3, start=1):
            max_len = max(len(str(h)), max(len(str(r.get(h))) for r in rows_kernel_step))
            ws3.column_dimensions[get_column_letter(j)].width = min(45, max(10, max_len + 2))
        ws3.freeze_panes = "A2"

    wb.save(excel_path)


def write_summary_excel_atomic(excel_path, rows_main, rows_kernel, rows_kernel_step):
    out_dir = os.path.dirname(os.path.abspath(excel_path))
    os.makedirs(out_dir, exist_ok=True)
    fd, tmp_path = tempfile.mkstemp(prefix="tmp_summary_", suffix=".xlsx", dir=out_dir)
    os.close(fd)
    try:
        write_summary_excel(tmp_path, rows_main, rows_kernel, rows_kernel_step)
        os.replace(tmp_path, excel_path)
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass


# ============================================================
# 11) Dynamic h candidates
# ============================================================
def build_dynamic_h_candidates(x_cur_norm, n_h=1000, h_min=0.01):
    x_cur_norm = float(x_cur_norm)
    h_min = float(h_min)
    if x_cur_norm <= h_min:
        return np.full(int(n_h), h_min, dtype=np.float64)
    return np.logspace(np.log10(h_min), np.log10(x_cur_norm), int(n_h))


# ============================================================
# 12) Per-step update of h and weights
#     coarse grid + local bounded optimization
# ============================================================
def _make_bandwidth_objective(
    X_events_norm,
    x_grid_norm,
    k_func,
    method,
    estimator_type="ordinary",
    v_grid_size=1000,
):
    X_events_norm = np.asarray(X_events_norm, dtype=np.float64)
    x_grid_norm = np.asarray(x_grid_norm, dtype=np.float64)

    def objective(h):
        h = float(h)
        _, llcv, lscv, campbell = compute_scores(
            X_events_norm,
            np.array([h], dtype=np.float64),
            x_grid_norm,
            k_func,
            estimator_type=estimator_type,
            v_grid_size=v_grid_size,
        )

        if method == "LLCV":
            return -float(llcv[0])
        elif method == "LSCV":
            return float(lscv[0])
        elif method == "Campbell":
            return float(campbell[0])
        else:
            raise ValueError(f"Unknown method: {method}")

    return objective


def _pick_h_by_coarse_grid_and_local_refine(
    X_events_norm,
    x_grid_norm,
    k_func,
    method,
    n_h,
    estimator_type="ordinary",
    h_min=0.01,
    v_grid_size=1000,
    coarse_ratio=0.1,
    coarse_min_points=61,
    local_neighbor_points=2,
    local_expand_factor=1.5,
):
    X_events_norm = np.asarray(X_events_norm, dtype=np.float64)
    x_cur_norm = float(np.max(X_events_norm)) if len(X_events_norm) else 0.0

    if x_cur_norm <= h_min:
        return float(h_min), {"mode": "degenerate", "coarse_idx": 0, "coarse_n_h": 1}

    global_left = float(h_min)
    global_right = float(x_cur_norm)

    coarse_n_h = int(max(coarse_min_points, np.ceil(n_h * coarse_ratio)))
    coarse_n_h = min(max(coarse_n_h, 3), max(int(n_h), 3))
    h_cands_coarse = build_dynamic_h_candidates(global_right, n_h=coarse_n_h, h_min=global_left)

    _, llcv, lscv, campbell = compute_scores(
        X_events_norm,
        h_cands_coarse,
        x_grid_norm,
        k_func,
        estimator_type=estimator_type,
        v_grid_size=v_grid_size,
    )

    if method == "LLCV":
        coarse_idx = int(np.argmax(llcv))
    elif method == "LSCV":
        coarse_idx = int(np.argmin(lscv))
    elif method == "Campbell":
        coarse_idx = int(np.argmin(campbell))
    else:
        raise ValueError(f"Unknown method: {method}")

    h0 = float(h_cands_coarse[coarse_idx])

    left_idx = max(0, coarse_idx - int(local_neighbor_points))
    right_idx = min(len(h_cands_coarse) - 1, coarse_idx + int(local_neighbor_points))
    local_left = float(h_cands_coarse[left_idx])
    local_right = float(h_cands_coarse[right_idx])

    local_left = max(global_left, min(local_left, h0 / float(local_expand_factor)))
    local_right = min(global_right, max(local_right, h0 * float(local_expand_factor)))

    if (not np.isfinite(local_left)) or (not np.isfinite(local_right)) or (local_right <= local_left):
        return h0, {
            "mode": "coarse_only",
            "coarse_idx": coarse_idx,
            "coarse_n_h": coarse_n_h,
            "local_left": local_left,
            "local_right": local_right,
        }

    objective = _make_bandwidth_objective(
        X_events_norm=X_events_norm,
        x_grid_norm=x_grid_norm,
        k_func=k_func,
        method=method,
        estimator_type=estimator_type,
        v_grid_size=v_grid_size,
    )

    try:
        res = minimize_scalar(
            objective,
            bounds=(local_left, local_right),
            method="bounded",
            options={"xatol": max(1e-6, 1e-3 * max(h0, h_min))},
        )
        if bool(getattr(res, "success", False)) and np.isfinite(res.x):
            h_best = float(res.x)
        else:
            h_best = h0
    except Exception:
        h_best = h0

    h_best = float(np.clip(h_best, global_left, global_right))
    return h_best, {
        "mode": "coarse_plus_local",
        "coarse_idx": coarse_idx,
        "coarse_n_h": coarse_n_h,
        "h0": h0,
        "local_left": local_left,
        "local_right": local_right,
    }



def update_kernel_configs_fixed_normalized_space(
    X_events_norm,
    kernel_map,
    method,
    n_h,
    x_grid_norm,
    estimator_type="ordinary",
    h_min=0.01,
    v_grid_size=1000,
):
    X_events_norm = np.asarray(X_events_norm, dtype=np.float64)
    x_cur_norm = float(np.max(X_events_norm)) if len(X_events_norm) else 0.0

    kernel_configs = []
    logliks = []

    for k_name, k_func in kernel_map.items():
        h_best, _ = _pick_h_by_coarse_grid_and_local_refine(
            X_events_norm=X_events_norm,
            x_grid_norm=x_grid_norm,
            k_func=k_func,
            method=method,
            n_h=n_h,
            estimator_type=estimator_type,
            h_min=h_min,
            v_grid_size=v_grid_size,
        )

        ell_m = loglik_mle_on_window(
            X_events_norm,
            float(h_best),
            k_func,
            x_grid_norm,
            x_cur=x_cur_norm,
            estimator_type=estimator_type,
            v_grid_size=v_grid_size,
        )

        kernel_configs.append({"name": k_name, "func": k_func, "h": float(h_best)})
        logliks.append(float(ell_m))

    return kernel_configs, np.asarray(logliks, dtype=np.float64)


def initialize_kernel_configs_and_fixed_weights_fixed_normalized_space(
    X_events_norm_init,
    kernel_map,
    method,
    n_h,
    x_grid_norm_init,
    estimator_type="ordinary",
    h_min=0.01,
    v_grid_size=1000,
):
    X_events_norm_init = np.asarray(X_events_norm_init, dtype=np.float64)

    kernel_configs_init, logliks_init = update_kernel_configs_fixed_normalized_space(
        X_events_norm=X_events_norm_init,
        kernel_map=kernel_map,
        method=method,
        n_h=n_h,
        x_grid_norm=x_grid_norm_init,
        estimator_type=estimator_type,
        h_min=h_min,
        v_grid_size=v_grid_size,
    )
    fixed_weights = max_likelihood_weights(logliks_init)
    return kernel_configs_init, fixed_weights, logliks_init


# ============================================================
# 13) Recursive prediction
# ============================================================
def recursive_predict_ensemble_one_time_rescale(
    X_obs_time,
    kernel_map,
    method,
    n_h,
    x_grid_base_unit,
    x_end_time,
    estimator_type="ordinary",
    max_steps=10000,
    verbose=False,
    h_min=0.01,
    v_grid_size=1000,
):
    X_obs_time = np.sort(np.asarray(X_obs_time, dtype=np.float64))
    t_cut = float(X_obs_time[-1])
    if t_cut <= 0.0:
        raise ValueError("t_cut must be positive.")

    X_events_norm = list(np.asarray(X_obs_time / t_cut, dtype=np.float64))
    X_events_norm.sort()

    x_end_norm = float(x_end_time / t_cut)

    preds_norm = []
    flags = []
    step_records = []

    x_cur_norm = float(X_events_norm[-1])

    grid_right_init = max(x_cur_norm, 1.0)
    n_grid_init = max(2000, int(np.ceil(2000 * grid_right_init)))
    x_grid_norm_init = np.linspace(0.0, grid_right_init, n_grid_init, dtype=np.float64)

    _, fixed_weights, init_logliks = initialize_kernel_configs_and_fixed_weights_fixed_normalized_space(
        X_events_norm_init=np.asarray(X_events_norm, dtype=np.float64),
        kernel_map=kernel_map,
        method=method,
        n_h=n_h,
        x_grid_norm_init=x_grid_norm_init,
        estimator_type=estimator_type,
        h_min=h_min,
        v_grid_size=v_grid_size,
    )

    for step in range(1, max_steps + 1):
        if x_cur_norm >= x_end_norm - 1e-12:
            break

        grid_right = max(x_cur_norm, 1.0)
        n_grid = max(2000, int(np.ceil(2000 * grid_right)))
        x_grid_norm = np.linspace(0.0, grid_right, n_grid, dtype=np.float64)

        kernel_configs, logliks = update_kernel_configs_fixed_normalized_space(
            X_events_norm=np.asarray(X_events_norm, dtype=np.float64),
            kernel_map=kernel_map,
            method=method,
            n_h=n_h,
            x_grid_norm=x_grid_norm,
            estimator_type=estimator_type,
            h_min=h_min,
            v_grid_size=v_grid_size,
        )

        n_used = int(len(X_events_norm))
        x_cur_time = float(x_cur_norm * t_cut)

        for cfg, ll, w in zip(kernel_configs, logliks, fixed_weights):
            step_records.append(
                {
                    "step": int(step),
                    "kernel": cfg["name"],
                    "h_best": float(cfg["h"]),
                    "ell_m": float(ll),
                    "weight": float(w),
                    "x_cur": float(x_cur_time),
                    "n_events_used": n_used,
                    "estimator_type": estimator_type,
                }
            )

        x_next_norm, ok = find_next_event_time_ensemble(
            x_cur_norm,
            np.asarray(X_events_norm, dtype=np.float64),
            kernel_configs,
            fixed_weights,
            x_end_norm,
            estimator_type=estimator_type,
            v_grid_size=v_grid_size,
        )

        if (not ok) or (x_next_norm <= x_cur_norm + 1e-10):
            break

        preds_norm.append(float(x_next_norm))
        flags.append(True)

        X_events_norm.append(float(x_next_norm))
        X_events_norm.sort()
        x_cur_norm = float(x_next_norm)

        if verbose and step % 10 == 0:
            print(
                f"[one-time-scale | fixed-weights][{estimator_type}] step={step}, "
                f"x_next_norm={x_next_norm:.6f}, ok={ok}",
                flush=True,
            )

        if np.isclose(x_cur_norm, x_end_norm) or x_cur_norm >= x_end_norm:
            break

    preds_time = np.asarray(preds_norm, dtype=np.float64) * t_cut
    return preds_time, np.asarray(flags, dtype=bool), step_records


# ============================================================
# 14) Worker warmup
# ============================================================
def init_worker():
    """
    Warm up Numba in each worker to reduce first-task latency spikes and
    avoid simultaneous first-compile contention.
    """
    try:
        faulthandler.enable()
    except Exception:
        pass

    try:
        x_grid = np.linspace(0.0, 1.0, 8, dtype=np.float64)
        X_events = np.array([0.1, 0.4, 0.9], dtype=np.float64)
        k_hat = np.array([1.0, 0.8, 1.2], dtype=np.float64)

        _ = _get_intensity_scalar_numba(x_grid, X_events, 0.2, 0)
        _ = _get_intensity_corrected_scalar_numba(x_grid, X_events, 0.2, 0, 0.0, 1.0, 16)
        _ = _loo_row_sums_scalar_numba(X_events, 0.2, 0)
        _ = _weighted_pred_intensity_from_khat_scalar_numba(x_grid, X_events, 0.2, k_hat, 0)
        _ = _cumulative_trapz_numba(x_grid, x_grid)
    except Exception as e:
        print(f"[Worker init warning] {e}", flush=True)


# ============================================================
# 15) Worker: one case = (data_name, frac)
# ============================================================
def _run_one_case_data_frac(
    data_name,
    raw_data,
    frac,
    methods,
    estimator_types,
    base_folder,
    n_h,
    x_grid_unit,
    h_min=0.01,
    v_grid_size=1000,
):
    setup_matplotlib_fonts()

    kernel_map = get_kernel_map()

    raw_sorted = np.sort(np.asarray(raw_data, dtype=np.float64))
    n_all = len(raw_sorted)

    sub_folder = os.path.join(base_folder, data_name)
    os.makedirs(sub_folder, exist_ok=True)

    frac_folder = os.path.join(sub_folder, f"OBS_{int(frac * 100)}pct")
    os.makedirs(frac_folder, exist_ok=True)

    m = max(3, int(np.ceil(n_all * frac)))
    raw_obs = raw_sorted[:m]
    raw_val = raw_sorted[m:]

    t_cut = float(raw_obs[-1])
    X_obs_time = np.sort(np.asarray(raw_obs, dtype=np.float64))
    X_all_time = np.asarray(raw_sorted, dtype=np.float64)
    X_val_time = np.asarray(raw_val, dtype=np.float64)
    x_end_time = float(X_all_time[-1])

    X_obs_norm = np.sort(np.asarray(raw_obs / t_cut, dtype=np.float64))

    emp_x_all = np.concatenate(([0.0], X_all_time, [x_end_time]))
    emp_y_all = np.concatenate(([0.0], np.arange(1, n_all + 1), [n_all]))

    summary_rows = []
    kernel_rows = []
    kernel_step_rows = []

    h_cands_init = build_dynamic_h_candidates(1.0, n_h=n_h, h_min=h_min)

    for estimator_type in estimator_types:
        # initial score plots
        for k_name, k_func in kernel_map.items():
            _, llcv, lscv, campbell = compute_scores(
                X_obs_norm,
                h_cands_init,
                x_grid_unit,
                k_func,
                estimator_type=estimator_type,
                v_grid_size=v_grid_size,
            )

            fig1, axes = plt.subplots(2, 2, figsize=(12, 10))
            fig1.suptitle(
                f"{data_name} | {k_name} | {estimator_type}\n"
                f"First {int(frac * 100)}% observed (n={len(X_obs_norm)}), normalized endpoint = 1 | Bandwidth Selection",
                fontsize=14,
            )

            axes[0, 1].semilogx(h_cands_init, llcv, "b")
            axes[0, 1].set_title("LLCV (max)")

            axes[1, 0].semilogx(h_cands_init, lscv, "orange")
            axes[1, 0].set_title("LSCV (min)")

            axes[1, 1].semilogx(h_cands_init, campbell, "g")
            axes[1, 1].set_title("Campbell (min)")

            axes[0, 0].axis("off")

            for ax in axes.flat:
                ax.grid(True, which="both", alpha=0.3)

            plt.tight_layout(rect=[0, 0.03, 1, 0.95])
            plt.savefig(
                os.path.join(
                    frac_folder,
                    f"{data_name}_{k_name}_{estimator_type}_1_Scores_OBS{int(frac * 100)}.png",
                )
            )
            plt.close(fig1)

        for method in methods:
            preds_time, ok_flags, step_records = recursive_predict_ensemble_one_time_rescale(
                X_obs_time=X_obs_time,
                kernel_map=kernel_map,
                method=method,
                n_h=n_h,
                x_grid_base_unit=x_grid_unit,
                x_end_time=x_end_time,
                estimator_type=estimator_type,
                verbose=False,
                h_min=h_min,
                v_grid_size=v_grid_size,
            )

            for r in step_records:
                kernel_step_rows.append(
                    {
                        "data": data_name,
                        "OBS_pct": int(frac * 100),
                        "estimator_type": estimator_type,
                        "method": method,
                        **r,
                    }
                )

            if len(step_records) > 0:
                last_step = max(rr["step"] for rr in step_records)
                last_pack = [rr for rr in step_records if rr["step"] == last_step]
                for rr in last_pack:
                    kernel_rows.append(
                        {
                            "data": data_name,
                            "OBS_pct": int(frac * 100),
                            "estimator_type": estimator_type,
                            "method": method,
                            "kernel": rr["kernel"],
                            "h_best": rr["h_best"],
                            "ell_m": rr["ell_m"],
                            "weight": rr["weight"],
                            "step_used": int(last_step),
                            "x_cur": rr["x_cur"],
                            "n_events_used": rr["n_events_used"],
                        }
                    )

            L = min(len(preds_time), len(X_val_time))
            if L > 0:
                err = preds_time[:L] - X_val_time[:L]
                mae = float(np.mean(np.abs(err)))
                rmse = float(np.sqrt(np.mean(err**2)))
            else:
                mae, rmse = np.nan, np.nan

            (
                PAE_MAE,
                PAE_RMSE,
                _,
                _,
                y_true_at_val,
                a_at_val,
                AE_curve,
            ) = compute_pae_metrics_piecewise_linear(X_obs_time, preds_time, X_val_time, x_end_time)

            ok_rate = float(np.mean(ok_flags)) if len(ok_flags) else np.nan
            tce = int(len(preds_time) - len(X_val_time))

            if (x_end_time - t_cut) <= 1e-12:
                ct_strict = np.nan
                ct_time = np.nan
            else:
                if len(preds_time) > 0:
                    last_pred = float(min(preds_time[-1], x_end_time))
                    ct_time = float((last_pred - t_cut) / (x_end_time - t_cut))

                    if np.any(ok_flags):
                        t_last_ok = float(preds_time[np.where(ok_flags)[0][-1]])
                    else:
                        t_last_ok = t_cut
                    ct_strict = float((t_last_ok - t_cut) / (x_end_time - t_cut))
                else:
                    ct_time = 0.0
                    ct_strict = 0.0

            metrics_line = (
                f"[DynamicHFixedW | {estimator_type} | {method} | one-time-rescale | h in [0.01, x_cur]] "
                f"pred_steps={len(preds_time)} | "
                f"Ct_strict={ct_strict:.4f} | Ct_time={ct_time:.4f} | TCE={tce:+d} | "
                f"ok_rate={ok_rate:.3f} | "
                f"MAE_vs_val={mae:.6f} | RMSE_vs_val={rmse:.6f} | "
                f"PAE_MAE={PAE_MAE:.6f} | PAE_RMSE={PAE_RMSE:.6f}"
            )

            metrics_path = os.path.join(
                frac_folder,
                f"{data_name}_DynamicHFixedW_{estimator_type}_{method}_Metrics_OBS{int(frac * 100)}.txt",
            )
            with open(metrics_path, "w", encoding="utf-8") as f:
                f.write(f"data={data_name}\n")
                f.write(f"OBS={int(frac * 100)}%\n")
                f.write(f"estimator_type={estimator_type}\n")
                f.write(f"method={method}\n")
                f.write(f"t_cut={t_cut}\n")
                f.write(f"x_end={x_end_time}\n\n")
                f.write(metrics_line + "\n")
                f.write("\n[PAE details on validation times]\n")
                f.write(f"PAE_MAE={PAE_MAE:.10f}\n")
                f.write(f"PAE_RMSE={PAE_RMSE:.10f}\n")
                if AE_curve is not None:
                    f.write(f"AE_curve_mean={np.mean(AE_curve):.10f}\n")
                    f.write(f"AE_curve_max={np.max(AE_curve):.10f}\n")

            pred_txt = os.path.join(
                frac_folder,
                f"{data_name}_DynamicHFixedW_{estimator_type}_{method}_RecursivePred_OBS{int(frac * 100)}.txt",
            )
            with open(pred_txt, "w", encoding="utf-8") as f:
                f.write(f"data={data_name}\n")
                f.write(f"OBS={int(frac * 100)}%\n")
                f.write(f"estimator_type={estimator_type}\n")
                f.write(f"method={method}\n")
                f.write(f"t_cut={t_cut}\n")
                f.write(f"x_end={x_end_time}\n\n")

                f.write("[Curve-based PAE Metrics]\n")
                f.write(f"PAE_MAE\t{PAE_MAE:.10f}\n")
                f.write(f"PAE_RMSE\t{PAE_RMSE:.10f}\n\n")

                f.write("[Per-step Kernel h_best & weights in fixed normalized space]\n")
                f.write("step\tx_cur\tn_events_used\tkernel\th_best\tell_m\tweight\n")
                for rr in step_records:
                    f.write(
                        f"{rr['step']}\t{rr['x_cur']:.10f}\t{rr['n_events_used']}\t"
                        f"{rr['kernel']}\t{rr['h_best']:.10f}\t{rr['ell_m']:.6f}\t{rr['weight']:.8f}\n"
                    )

                if X_val_time is not None and len(X_val_time) > 0 and AE_curve is not None:
                    f.write("\n[PAE Curve on Validation Times]\n")
                    f.write("idx\tX_val\ty_true_at_val\ta_pred_piecewise_linear\tAE\n")
                    for i in range(len(X_val_time)):
                        f.write(
                            f"{i + 1}\t{X_val_time[i]:.10f}\t{y_true_at_val[i]:.10f}\t"
                            f"{a_at_val[i]:.10f}\t{AE_curve[i]:.10f}\n"
                        )

                f.write("\n[Recursive Pred Sequence]\n")
                f.write("step\tpred_x\tok_flag\ttrue_val_x(if any)\n")
                for i, px in enumerate(preds_time):
                    tv = X_val_time[i] if i < len(X_val_time) else np.nan
                    ok = int(ok_flags[i]) if i < len(ok_flags) else 0
                    f.write(f"{i + 1}\t{px:.10f}\t{ok}\t{tv:.10f}\n")

            # Plot cumulative curve + last-step weights pie chart
            fig = plt.figure(figsize=(14, 6))
            gs = fig.add_gridspec(1, 3)
            ax1 = fig.add_subplot(gs[0, :2])

            ax1.step(
                emp_x_all,
                emp_y_all,
                where="post",
                color="black",
                lw=2,
                label="Empirical Cumulative Count (All Data)",
                zorder=5,
            )
            ax1.axvline(t_cut, linestyle="--", alpha=0.6)
            ax1.scatter(X_all_time, np.arange(1, n_all + 1), s=10, alpha=0.25)

            X_pred_all = np.concatenate([X_obs_time, preds_time])
            X_pred_all.sort()
            y_pred = np.arange(1, len(X_pred_all) + 1)

            ax1.step(
                np.concatenate(([0.0], X_pred_all, [x_end_time])),
                np.concatenate(([0.0], y_pred, [y_pred[-1] if len(y_pred) else 0.0])),
                where="post",
                lw=2,
                alpha=0.85,
                label=(
                    f"DynamicHFixedW({estimator_type}, {method}) "
                    f"steps={len(preds_time)} | MAE={mae:.4f} | PAE_MAE={PAE_MAE:.4f}"
                ),
            )

            x_knots_plot = np.concatenate(([0.0], X_pred_all, [x_end_time]))
            y_knots_plot = np.concatenate(
                ([0.0], np.arange(1, len(X_pred_all) + 1), [len(X_pred_all)])
            )
            ax1.plot(
                x_knots_plot,
                y_knots_plot,
                alpha=0.5,
                linewidth=1.5,
                linestyle="-",
                label="Predicted Cumulative Count (Piecewise Linear)",
            )

            ax1.set_title(
                f"{data_name} | DynamicHFixedW({estimator_type}) | {method}\n"
                f"OBS={int(frac * 100)}% | Reference = t_cut (original time unit)"
            )
            ax1.set_xlabel("Time t (original unit)")
            ax1.set_ylabel("Cumulative Number of Failures")
            ax1.grid(True, alpha=0.2)
            ax1.legend(fontsize="small")

            ax2 = fig.add_subplot(gs[0, 2])
            if len(step_records) > 0:
                last_step = max(rr["step"] for rr in step_records)
                last_pack = [rr for rr in step_records if rr["step"] == last_step]
                labels = [rr["kernel"] for rr in last_pack]
                weights_last = np.array([rr["weight"] for rr in last_pack], dtype=np.float64)

                plot_w, plot_l = [], []
                for w, l in zip(weights_last, labels):
                    if w > 0.01:
                        plot_w.append(w)
                        plot_l.append(l)
                if len(plot_w) == 0:
                    plot_w = list(weights_last)
                    plot_l = list(labels)

                ax2.pie(
                    plot_w,
                    labels=plot_l,
                    autopct="%1.1f%%",
                    startangle=140,
                    colors=plt.cm.Pastel1.colors,
                )
                ax2.set_title(f"Kernel Weights (fixed weights; shown at step = {last_step})")
            else:
                ax2.text(0.5, 0.5, "No weights", ha="center", va="center")
                ax2.set_axis_off()

            plt.tight_layout()
            plt.savefig(
                os.path.join(
                    frac_folder,
                    f"{data_name}_DynamicHFixedW_{estimator_type}_{method}_Cum_OBS{int(frac * 100)}.png",
                )
            )
            plt.close(fig)

            summary_rows.append(
                {
                    "data": data_name,
                    "OBS_pct": int(frac * 100),
                    "estimator_type": estimator_type,
                    "method": method,
                    "n_all": int(n_all),
                    "obs_n": int(len(X_obs_time)),
                    "val_n": int(len(X_val_time)),
                    "t_cut": float(t_cut),
                    "x_end": float(x_end_time),
                    "pred_steps": int(len(preds_time)),
                    "Ct_strict": float(ct_strict) if not np.isnan(ct_strict) else None,
                    "Ct_time": float(ct_time) if not np.isnan(ct_time) else None,
                    "TCE": int(tce),
                    "ok_rate": float(ok_rate) if not np.isnan(ok_rate) else None,
                    "MAE_vs_val": float(mae) if not np.isnan(mae) else None,
                    "RMSE_vs_val": float(rmse) if not np.isnan(rmse) else None,
                    "PAE_MAE": float(PAE_MAE) if not np.isnan(PAE_MAE) else None,
                    "PAE_RMSE": float(PAE_RMSE) if not np.isnan(PAE_RMSE) else None,
                }
            )

    return summary_rows, kernel_rows, kernel_step_rows


# ============================================================
# 16) Parallel runner
# ============================================================
def run_batch_analysis_ensemble_dynamic_h_fixed_w_parallel(
    dataset_dict,
    base_folder="ensemble_dynamic_h_fixed_w",
    fractions=(0.2, 0.5, 0.8),
    estimator_types=("ordinary",),
    methods=("LLCV", "LSCV", "Campbell"),
    summary_excel_name="ALL_RESULTS_SUMMARY1.xlsx",
    max_workers=None,
    mp_start_method="spawn",
    warmup_workers=True,
):
    setup_matplotlib_fonts()
    os.makedirs(base_folder, exist_ok=True)

    summary_rows = []
    kernel_rows = []
    kernel_step_rows = []
    failed_rows = []

    n_h = 1000
    x_grid_unit = np.linspace(0.0, 1.0, 2000, dtype=np.float64)

    tasks = []
    for data_name, raw_data in dataset_dict.items():
        for frac in fractions:
            tasks.append((data_name, np.asarray(raw_data, dtype=np.float64), float(frac)))

    excel_path = os.path.join(base_folder, summary_excel_name)

    max_workers = recommend_max_workers(max_workers)

    print(
        f"[Parallel] total_cases={len(tasks)} | max_workers={max_workers} | "
        f"start_method={mp_start_method} | numba_warmup={warmup_workers}",
        flush=True,
    )
    print(f"[Excel will be updated per completed case] -> {os.path.abspath(excel_path)}", flush=True)

    mp_context = mp.get_context(mp_start_method)
    initializer = init_worker if warmup_workers else None

    with ProcessPoolExecutor(
        max_workers=max_workers,
        mp_context=mp_context,
        initializer=initializer,
    ) as ex:
        future_map = {}
        for data_name, raw_data, frac in tasks:
            fut = ex.submit(
                _run_one_case_data_frac,
                data_name,
                raw_data,
                frac,
                methods,
                estimator_types,
                base_folder,
                n_h,
                x_grid_unit,
                0.01,
                1000,
            )
            future_map[fut] = {
                "data_name": data_name,
                "frac": frac,
            }

        done_cases = 0
        total_cases = len(future_map)

        for fut in as_completed(future_map):
            spec = future_map[fut]
            try:
                case_summary, case_kernel, case_kernel_step = fut.result()
                summary_rows.extend(case_summary)
                kernel_rows.extend(case_kernel)
                kernel_step_rows.extend(case_kernel_step)

                done_cases += 1
                write_summary_excel_atomic(excel_path, summary_rows, kernel_rows, kernel_step_rows)
                print(
                    f"[Case done] {done_cases}/{total_cases} | "
                    f"data={spec['data_name']} | OBS={int(spec['frac'] * 100)}% -> Excel updated",
                    flush=True,
                )
            except Exception as e:
                err_text = "".join(traceback.format_exception(type(e), e, e.__traceback__))
                failed_rows.append(
                    {
                        "data": spec["data_name"],
                        "OBS_pct": int(spec["frac"] * 100),
                        "error": str(e),
                    }
                )
                print(
                    f"[Case failed] data={spec['data_name']} | OBS={int(spec['frac'] * 100)}%\n{err_text}",
                    flush=True,
                )
            finally:
                gc.collect()

    if failed_rows:
        failed_csv = os.path.join(base_folder, "FAILED_CASES.csv")
        pd.DataFrame(failed_rows).to_csv(failed_csv, index=False, encoding="utf-8-sig")
        print(f"[Failed cases saved] {os.path.abspath(failed_csv)}", flush=True)

    print(f"\n[All completed] Excel summary saved at: {os.path.abspath(excel_path)}", flush=True)


# ============================================================
# 17) Data loading
# ============================================================
def load_datasets_from_data_folder(data_folder="DATA"):
    data_dir = os.path.join(os.getcwd(), data_folder)
    dataset_dict = {}

    if not os.path.exists(data_dir):
        print(f"Error: folder not found: {data_dir}")
        return dataset_dict

    file_list = glob.glob(os.path.join(data_dir, "*.xlsx"))

    for file_path in file_list:
        file_name = os.path.basename(file_path)
        dataset_name = os.path.splitext(file_name)[0]

        try:
            df = pd.read_excel(file_path, header=None)
            raw_data = df.iloc[:, 0].dropna().values

            if len(raw_data) > 0:
                dataset_dict[dataset_name] = np.sort(np.asarray(raw_data, dtype=np.float64))
                print(
                    f"Successfully loaded dataset: {dataset_name} | Number of events: {len(raw_data)}",
                    flush=True,
                )
            else:
                print(f"Skipped file {file_name}: empty data", flush=True)

        except Exception as e:
            print(f"Skipped file {file_name}: failed to read ({e})", flush=True)

    return dataset_dict


# ============================================================
# 18) Main entry
# ============================================================
if __name__ == "__main__":
    freeze_support()

    safe_send_notice_mail(
        "Started",
        "The task has started!",
        "d252433@hiroshima-u.ac.jp",
    )

    my_datasets = load_datasets_from_data_folder("DATA")

    if not my_datasets:
        print(
            "No valid dataset detected. Program exits. Please check whether the DATA folder contains .xlsx files.",
            flush=True,
        )
    else:
        out_dir = "20260403_DYNAMIC_H_AND_FIXED_WEIGHT_Numbarized_ServerStable"

        run_batch_analysis_ensemble_dynamic_h_fixed_w_parallel(
            my_datasets,
            base_folder=out_dir,
            fractions=(0.2, 0.5, 0.8),
            estimator_types=("ordinary",),
            methods=("LLCV", "LSCV", "Campbell"),
            summary_excel_name="ALL_RESULTS_SUMMARY1.xlsx",
            max_workers=30,          # recommend_max_workers will cap safely if needed
            mp_start_method="spawn", # safest on servers for JIT + plotting + multiprocessing
            warmup_workers=True,     # avoid first-compile spikes / cache races
        )

        safe_send_notice_mail(
            "Finished",
            "The task has been completed!",
            "d252433@hiroshima-u.ac.jp",
        )

        print(f"\n[All tasks completed] Results saved in: {os.path.abspath(out_dir)}", flush=True)
