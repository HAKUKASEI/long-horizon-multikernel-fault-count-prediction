"""Reconstruct and validate the paper's four OSS event-time arrays."""

import argparse
import csv
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


@dataclass(frozen=True)
class DatasetSpec:
    dataset_id: str
    project: str
    count: int
    maximum: int
    origin_issue: int
    without_bug_label: int


SPECS = (
    DatasetSpec("OSS_1", "ChatGPT", 278, 10437067, 1, 0),
    DatasetSpec("OSS_2", "MobX", 278, 220139318, 191, 6),
    DatasetSpec("OSS_3", "XGBoost", 183, 144869150, 3579, 17),
    DatasetSpec("OSS_4", "uni-app", 290, 127116130, 103, 2),
)


@dataclass(frozen=True)
class Record:
    issue_number: int
    created_at: datetime
    categories: str


def parse_utc(value):
    if not isinstance(value, str):
        raise ValueError("Original Created At Text must be a UTC timestamp string")
    try:
        timestamp = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as error:
        raise ValueError(f"Invalid UTC timestamp: {value!r}") from error
    if timestamp.utcoffset() is None or timestamp.utcoffset().total_seconds() != 0:
        raise ValueError(f"Timestamp must explicitly use UTC: {value!r}")
    if timestamp.microsecond:
        raise ValueError(f"Expected whole-second timestamp: {value!r}")
    return timestamp.astimezone(timezone.utc)


def build_sequence(records):
    if len(records) < 2:
        raise ValueError("At least an origin record and one subsequent event are required")
    ordered = sorted(records, key=lambda record: (record.created_at, record.issue_number))
    if len({record.issue_number for record in ordered}) != len(ordered):
        raise ValueError("Duplicate issue numbers within a dataset")
    origin = ordered[0]
    times = [int((record.created_at - origin.created_at).total_seconds()) for record in ordered[1:]]
    if any(value <= 0 for value in times):
        raise ValueError("Model times must be positive; the earliest record only defines the origin")
    return ordered, times


def rows_by_header(sheet):
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header or len(set(header)) != len(header):
        raise ValueError(f"Missing or duplicate headers in {sheet.title}")
    return [dict(zip(header, row)) for row in rows if any(value is not None for value in row)]


def read_array(sheet):
    values = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if len(row) != 1:
            raise ValueError(f"{sheet.title}: expected exactly one column")
        value = row[0]
        if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
            raise ValueError(f"{sheet.title}, row {index}: expected a headerless integer time value")
        values.append(int(value))
    return values


def validate(input_path, array_dir):
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        raw = rows_by_header(workbook["Raw_Issue_Records"])
        metadata_rows = rows_by_header(workbook["Dataset_Metadata"])
        metadata = {row["Dataset ID"]: row for row in metadata_rows}
        ids = {spec.dataset_id for spec in SPECS}
        if len(metadata_rows) != len(ids) or set(metadata) != ids:
            raise ValueError("Metadata must contain each of the four paper dataset IDs exactly once")
        grouped = {dataset_id: [] for dataset_id in ids}
        for row in raw:
            dataset_id = row["Dataset ID"]
            if dataset_id not in grouped:
                raise ValueError(f"Unexpected dataset ID: {dataset_id!r}")
            issue = row["Issue Number"]
            if isinstance(issue, bool) or not isinstance(issue, int) or issue <= 0:
                raise ValueError(f"Invalid issue number: {issue!r}")
            labels = row["Categories"] or ""
            if row["Has Bug Label"] != ("bug" in labels.casefold()):
                raise ValueError(f"{dataset_id}, issue {issue}: inconsistent label flag")
            grouped[dataset_id].append(Record(issue, parse_utc(row["Original Created At Text"]), labels))
        results = []
        for spec in SPECS:
            ordered, times = build_sequence(grouped[spec.dataset_id])
            without_bug = sum("bug" not in record.categories.casefold() for record in ordered[1:])
            controls = (len(times), times[-1], ordered[0].issue_number, without_bug)
            expected = (spec.count, spec.maximum, spec.origin_issue, spec.without_bug_label)
            if controls != expected:
                raise ValueError(f"{spec.dataset_id}: paper controls differ: {controls} != {expected}")
            meta = metadata[spec.dataset_id]
            expected_metadata = {
                "Project": spec.project, "Model Events": spec.count, "Start (sec)": 0,
                "End (sec)": spec.maximum, "Time Unit": "Calendar sec", "Raw Records": len(ordered),
                "Origin Issue": spec.origin_issue, "First Model Event (sec)": times[0],
                "Model Records Without Bug Label": without_bug,
            }
            for field, expected_value in expected_metadata.items():
                if meta[field] != expected_value:
                    raise ValueError(f"{spec.dataset_id}: metadata mismatch for {field}")
            if read_array(workbook[spec.dataset_id]) != times:
                raise ValueError(f"{spec.dataset_id}: combined workbook array differs from reconstructed timestamps")
            array_path = array_dir / f"{spec.dataset_id}.xlsx"
            standalone = load_workbook(array_path, read_only=True, data_only=True)
            try:
                if len(standalone.sheetnames) != 1 or read_array(standalone.active) != times:
                    raise ValueError(f"{array_path.name}: experiment input differs from reconstructed timestamps")
            finally:
                standalone.close()
            results.append({
                "dataset_id": spec.dataset_id, "project": spec.project, "model_events": len(times),
                "raw_records": len(ordered), "origin_utc": ordered[0].created_at.strftime("%Y-%m-%dT%H:%M:%SZ"),
                "first_seconds": times[0], "maximum_seconds": times[-1],
                "model_records_without_bug_label": without_bug, "times": times,
            })
        return results
    finally:
        workbook.close()


def write_csvs(results, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    for result in results:
        destination = output_dir / f"{result['dataset_id']}_times.csv"
        with destination.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            writer.writerows((time,) for time in result["times"])
    fields = [key for key in results[0] if key != "times"]
    with (output_dir / "dataset_summary.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows({key: result[key] for key in fields} for result in results)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=ROOT / "data/JSS_OSS_time_datasets_OSS1_OSS4.xlsx")
    parser.add_argument("--array-dir", type=Path, default=ROOT / "data/OSS")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data/processed_time")
    parser.add_argument("--check-only", action="store_true", help="Validate without writing CSV outputs")
    args = parser.parse_args(argv)
    try:
        results = validate(args.input, args.array_dir)
        for result in results:
            print(f"[OK] {result['dataset_id']} {result['project']}: {result['model_events']} events, "
                  f"first={result['first_seconds']}, max={result['maximum_seconds']} sec; all values match")
        if not args.check_only:
            write_csvs(results, args.output_dir)
            print(f"CSV outputs: {args.output_dir}")
    except (ValueError, KeyError, OSError, TypeError, OverflowError) as error:
        print(f"[ERROR] {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
